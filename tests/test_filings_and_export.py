from __future__ import annotations

import asyncio
import io
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook
from rich.console import Console
from textual.widgets import Static

from ajax_terminal.app import AjaxTerminalApp
from ajax_terminal.models.filing import Filing, FilingCollection
from ajax_terminal.models.quote import (
    DataQuality,
    FinancialPeriod,
    FinancialStatements,
    StatementType,
)
from ajax_terminal.providers.sec import (
    SECFilingsProvider,
    _PresentationItem,
    _StatementPresentation,
    _parse_statement_report,
    _statement_report_names,
    normalize_sec_companyfacts,
)
from ajax_terminal.providers.mock import MockMacroProvider, MockNewsProvider
from ajax_terminal.services.excel_export_service import ExcelExportService
from ajax_terminal.services.excel_export_service import FinancialExportResult
from ajax_terminal.services.market_service import MarketService
from ajax_terminal.storage.cache import SQLiteCache, WatchlistStore
from ajax_terminal.ui.commands import parse_command


def test_sec_provider_resolves_ticker_and_builds_official_document_url() -> None:
    provider = SECFilingsProvider()

    def fake_json(url: str):
        if url == provider.TICKERS_URL:
            return {"0": {"cik_str": 320193, "ticker": "AAPL", "title": "Apple Inc."}}
        return {
            "name": "Apple Inc.",
            "filings": {
                "recent": {
                    "form": ["10-Q", "8-K", "10-K"],
                    "filingDate": ["2026-08-01", "2026-07-10", "2025-10-31"],
                    "reportDate": ["2026-06-27", "2026-07-10", "2025-09-27"],
                    "accessionNumber": [
                        "0000320193-26-000070",
                        "0000320193-26-000060",
                        "0000320193-25-000079",
                    ],
                    "primaryDocument": ["aapl-20260627.htm", "aapl-8k.htm", "aapl-20250927.htm"],
                    "primaryDocDescription": ["10-Q", "8-K", "10-K"],
                }
            },
        }

    provider._get_json = fake_json  # type: ignore[method-assign]
    result = asyncio.run(provider.filings("AAPL", ("10-K", "10-Q")))

    assert result.quality == DataQuality.REALTIME
    assert [item.form for item in result.filings] == ["10-Q", "10-K"]
    assert result.filings[0].document_url == (
        "https://www.sec.gov/Archives/edgar/data/320193/"
        "000032019326000070/aapl-20260627.htm"
    )


def test_sec_provider_does_not_map_foreign_suffix_to_us_adr() -> None:
    provider = SECFilingsProvider()
    provider._ticker_index = {"SAN": ("0000891478", "Banco Santander, S.A.")}

    result = asyncio.run(provider.filings("SAN.MC"))

    assert result.quality == DataQuality.UNAVAILABLE
    assert result.filings == []


class StatementMarketService:
    def __init__(self, quality: DataQuality = DataQuality.DELAYED) -> None:
        self.quality = quality

    async def financial_statements(
        self,
        symbol: str,
        statement_type: StatementType,
    ) -> FinancialStatements:
        return FinancialStatements(
            symbol=symbol,
            name="Apple Inc.",
            statement_type=statement_type,
            annual=[FinancialPeriod("2025", date(2025, 9, 27), {"Revenue": 416_000_000_000.0, "CostOfRevenue": 221_000_000_000.0})],
            quarterly=[FinancialPeriod("Q1 2026", date(2025, 12, 27), {"Revenue": 94_000_000_000.0, "CostOfRevenue": 50_000_000_000.0})],
            currency="USD",
            provider="TEST PROVIDER",
            quality=self.quality,
        )


def test_excel_export_creates_annual_and_quarterly_statement_sheets(tmp_path) -> None:
    service = ExcelExportService(StatementMarketService(), tmp_path)  # type: ignore[arg-type]

    result = asyncio.run(service.export_financials("AAPL"))
    workbook = load_workbook(result.path, data_only=True)

    assert result.path.exists()
    assert set(result.sheet_names) == {
        "Income Annual",
        "Income Quarterly",
        "Balance Annual",
        "Balance Quarterly",
        "Cash Flow Annual",
        "Cash Flow Quarterly",
    }
    assert workbook["Income Annual"]["A11"].value == "Revenue"
    assert workbook["Income Annual"]["B11"].value == "USD M"
    assert workbook["Income Annual"]["C11"].value == 416_000.0
    assert workbook["Income Annual"]["C6"].value == "FY 2025"
    assert workbook["Income Annual"]["C7"].value.date() == date(2025, 9, 27)
    assert workbook["Income Annual"]["A8"].value == "Source Form"
    assert workbook["Income Annual"]["A12"].value == "Cost Of Revenue"
    assert workbook["Income Annual"]["A12"].font.color.rgb == "00C00000"
    assert workbook["Income Annual"]["C12"].font.color.rgb == "00C00000"
    assert workbook["Income Quarterly"]["C6"].value == "Q1 2026"
    assert workbook["Metadata"]["B2"].value == "AAPL"


def _sec_fact(
    value: float,
    start: str | None,
    end: str,
    filed: str,
    fiscal_year: int,
    fiscal_period: str,
    form: str,
    accession: str,
) -> dict[str, object]:
    row: dict[str, object] = {
        "end": end,
        "val": value,
        "filed": filed,
        "fy": fiscal_year,
        "fp": fiscal_period,
        "form": form,
        "accn": accession,
    }
    if start is not None:
        row["start"] = start
    return row


def test_sec_companyfacts_uses_original_report_periods_and_detailed_lines() -> None:
    annual = _sec_fact(400.0, "2024-09-29", "2025-09-27", "2025-10-31", 2025, "FY", "10-K", "annual")
    late_comparative = _sec_fact(400.0, "2024-09-29", "2025-09-27", "2026-10-30", 2026, "FY", "10-K", "comparison")
    q2_ytd = _sec_fact(220.0, "2024-09-29", "2025-03-29", "2025-05-02", 2025, "Q2", "10-Q", "q2")
    q2 = _sec_fact(95.0, "2024-12-29", "2025-03-29", "2025-05-02", 2025, "Q2", "10-Q", "q2")
    payload = {
        "entityName": "Example Corp",
        "facts": {
            "us-gaap": {
                "RevenueFromContractWithCustomerExcludingAssessedTax": {
                    "units": {"USD": [annual, late_comparative, q2_ytd, q2]}
                },
                "GrossProfit": {
                    "units": {
                        "USD": [
                            _sec_fact(180.0, "2024-09-29", "2025-09-27", "2025-10-31", 2025, "FY", "10-K", "annual")
                        ]
                    }
                },
            }
        },
    }

    statements = normalize_sec_companyfacts("TEST", "Example Corp", payload, StatementType.INCOME)

    assert statements.provider == "SEC EDGAR XBRL"
    assert statements.annual[0].period == "FY 2025"
    assert statements.annual[0].accession_number == "annual"
    assert statements.annual[0].values == {"totalRevenue": 400.0, "grossProfit": 180.0}
    assert statements.quarterly[0].period == "Q2 2025"
    assert statements.quarterly[0].values["totalRevenue"] == 95.0


def test_sec_filed_statement_presentation_adds_unmapped_primary_lines_and_labels() -> None:
    payload = {
        "entityName": "Example Corp",
        "facts": {
            "us-gaap": {
                "SellingAndMarketingExpense": {
                    "label": "Selling and Marketing Expense",
                    "units": {
                        "USD": [
                            _sec_fact(
                                25.0,
                                "2024-01-01",
                                "2024-12-31",
                                "2025-02-10",
                                2024,
                                "FY",
                                "10-K",
                                "annual",
                            )
                        ]
                    },
                }
            }
        },
    }
    presentation = _StatementPresentation(
        (
            _PresentationItem(
                "us-gaap",
                "SellingAndMarketingExpense",
                "Sales and marketing",
                "OPERATING EXPENSES",
            ),
        ),
        "https://www.sec.gov/Archives/example/R3.htm",
    )

    statements = normalize_sec_companyfacts(
        "TEST",
        "Example Corp",
        payload,
        StatementType.INCOME,
        presentation=presentation,
    )

    assert statements.annual[0].values["sellingAndMarketingExpense"] == 25.0
    assert statements.metric_labels["sellingAndMarketingExpense"] == "Sales and marketing"
    assert statements.metric_sections["sellingAndMarketingExpense"] == "OPERATING EXPENSES"
    assert statements.source_url.endswith("R3.htm")


def test_sec_filing_summary_and_report_parser_select_primary_rows_only() -> None:
    summary = """<FilingSummary><MyReports>
    <Report><HtmlFileName>R3.htm</HtmlFileName><ShortName>CONSOLIDATED STATEMENTS OF OPERATIONS</ShortName><MenuCategory>Statements</MenuCategory></Report>
    <Report><HtmlFileName>R4.htm</HtmlFileName><ShortName>CONSOLIDATED STATEMENTS OF COMPREHENSIVE INCOME</ShortName><MenuCategory>Statements</MenuCategory></Report>
    </MyReports></FilingSummary>"""
    report = """<table>
    <tr class="re"><td class="pl"><a onclick="Show.showAR(this, 'defref_us-gaap_OperatingExpensesAbstract', window)"><strong>Operating expenses:</strong></a></td></tr>
    <tr class="ro"><td class="pl"><a onclick="Show.showAR(this, 'defref_us-gaap_SellingAndMarketingExpense', window)">Sales and marketing</a></td><td class="nump">25</td></tr>
    <tr class="rh"><td class="pl"><a onclick="Show.showAR(this, 'defref_srt_ProductOrServiceAxis=us-gaap_ProductMember', window)">Products</a></td></tr>
    <tr class="ro"><td class="pl"><a onclick="Show.showAR(this, 'defref_us-gaap_RevenueFromContractWithCustomerExcludingAssessedTax', window)">Segment revenue</a></td></tr>
    </table>"""

    assert _statement_report_names(summary, StatementType.INCOME) == ("R3.htm", "R4.htm")
    items = _parse_statement_report(report, StatementType.INCOME)
    assert [(item.concept, item.label, item.section) for item in items] == [
        ("SellingAndMarketingExpense", "Sales and marketing", "OPERATING EXPENSES")
    ]


def test_sec_cash_flow_derives_standalone_quarters_from_reported_ytd_values() -> None:
    def rows(values: tuple[float, float, float, float]) -> list[dict[str, object]]:
        q1, half, nine_months, annual = values
        return [
            _sec_fact(q1, "2024-09-29", "2024-12-28", "2025-01-31", 2025, "Q1", "10-Q", "q1"),
            _sec_fact(half, "2024-09-29", "2025-03-29", "2025-05-02", 2025, "Q2", "10-Q", "q2"),
            _sec_fact(nine_months, "2024-09-29", "2025-06-28", "2025-08-01", 2025, "Q3", "10-Q", "q3"),
            _sec_fact(annual, "2024-09-29", "2025-09-27", "2025-10-31", 2025, "FY", "10-K", "fy"),
        ]

    payload = {
        "entityName": "Example Corp",
        "facts": {
            "us-gaap": {
                "NetCashProvidedByUsedInOperatingActivities": {"units": {"USD": rows((50.0, 80.0, 120.0, 150.0))}},
                "PaymentsToAcquirePropertyPlantAndEquipment": {"units": {"USD": rows((5.0, 8.0, 12.0, 16.0))}},
            }
        },
    }

    statements = normalize_sec_companyfacts("TEST", "Example Corp", payload, StatementType.CASH_FLOW)
    by_period = {period.period: period for period in statements.quarterly}

    assert by_period["Q2 2025"].derived is True
    assert by_period["Q2 2025"].values["operatingCashFlow"] == 30.0
    assert by_period["Q2 2025"].values["capitalExpenditure"] == -3.0
    assert by_period["Q2 2025"].values["freeCashFlowCalculated"] == 27.0
    assert by_period["Q4 2025"].values["operatingCashFlow"] == 30.0


def test_excel_export_refuses_mock_statements(tmp_path) -> None:
    service = ExcelExportService(StatementMarketService(DataQuality.MOCK), tmp_path)  # type: ignore[arg-type]

    with pytest.raises(ValueError, match="No real financial statements"):
        asyncio.run(service.export_financials("AAPL"))


class AppFilingsService:
    async def filings(self, symbol: str, forms=("10-K", "10-Q"), limit: int = 20):
        return FilingCollection(
            symbol=symbol,
            company_name="Apple Inc.",
            cik="0000320193",
            filings=[
                Filing(
                    symbol=symbol,
                    company_name="Apple Inc.",
                    cik="0000320193",
                    form=forms[0],
                    filing_date=date(2025, 10, 31),
                    report_date=date(2025, 9, 27),
                    accession_number="0000320193-25-000079",
                    primary_document="aapl-20250927.htm",
                    description="Annual report",
                )
            ],
            quality=DataQuality.REALTIME,
        )


class AppExportService:
    async def export_financials(self, symbol: str) -> FinancialExportResult:
        return FinancialExportResult(
            symbol=symbol,
            path=Path("exports/AAPL_financials.xlsx").resolve(),
            sheet_names=("Income Annual", "Cash Flow Quarterly"),
            skipped=("Balance Quarterly",),
        )


def test_filings_and_export_commands_render_in_terminal(tmp_path) -> None:
    async def run() -> None:
        database = tmp_path / "filings-ui.sqlite3"
        market = MarketService(
            cache=SQLiteCache(database),
            watchlists=WatchlistStore(database),
            market_providers=[],
        )
        app = AjaxTerminalApp(
            market_service=market,
            filings_service=AppFilingsService(),  # type: ignore[arg-type]
            export_service=AppExportService(),  # type: ignore[arg-type]
        )
        app.news_service.providers = [MockNewsProvider()]
        app.macro_service.providers = [MockMacroProvider()]
        async with app.run_test(size=(160, 45)) as pilot:
            await pilot.pause()
            await app.execute_command(parse_command("10K AAPL"))
            console = Console(width=160, record=True, file=io.StringIO())
            console.print(app.query_one("#main", Static).content)
            text = console.export_text()
            assert "OPEN SEC" in text
            assert "0000320193-25-000079" in text
            filing_header = str(app.query_one("#instrument-function", Static).content)
            assert "AAPL Equity" in filing_header
            assert "ANNUAL FILING" in filing_header

            await app.execute_command(parse_command("XLS AAPL"))
            console = Console(width=160, record=True, file=io.StringIO())
            console.print(app.query_one("#main", Static).content)
            text = console.export_text()
            assert "EXPORT COMPLETE" in text
            assert "AAPL_financials.xlsx" in text

    asyncio.run(run())
