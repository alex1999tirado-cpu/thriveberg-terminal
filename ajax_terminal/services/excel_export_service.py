from __future__ import annotations

import asyncio
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ajax_terminal.models.quote import DataQuality, FinancialPeriod, FinancialStatements, StatementType
from ajax_terminal.services.market_service import MarketService
from ajax_terminal.utils.financials import is_expense_or_outflow_metric, ordered_statement_metrics


@dataclass(frozen=True, slots=True)
class FinancialExportResult:
    symbol: str
    path: Path
    sheet_names: tuple[str, ...]
    skipped: tuple[str, ...]


class ExcelExportService:
    def __init__(self, market_service: MarketService, export_dir: Path | str = "exports") -> None:
        self.market_service = market_service
        self.export_dir = Path(export_dir)

    async def export_financials(self, symbol: str) -> FinancialExportResult:
        clean = symbol.strip().upper()
        statements = await asyncio.gather(
            self.market_service.financial_statements(clean, StatementType.INCOME),
            self.market_service.financial_statements(clean, StatementType.BALANCE_SHEET),
            self.market_service.financial_statements(clean, StatementType.CASH_FLOW),
        )
        real = [
            item
            for item in statements
            if item.quality not in {DataQuality.MOCK, DataQuality.UNAVAILABLE}
            and (item.annual or item.quarterly)
        ]
        if not real:
            raise ValueError("No real financial statements are available to export")
        self.export_dir.mkdir(parents=True, exist_ok=True)
        safe_symbol = re.sub(r"[^A-Z0-9._-]+", "_", clean)
        timestamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
        path = (self.export_dir / f"{safe_symbol}_financials_{timestamp}.xlsx").resolve()
        return await asyncio.to_thread(_write_workbook, clean, real, path)


def _write_workbook(
    symbol: str,
    statements: list[FinancialStatements],
    path: Path,
) -> FinancialExportResult:
    workbook = Workbook()
    workbook.remove(workbook.active)
    metadata = workbook.create_sheet("Metadata")
    metadata_rows = [
        ("THRIVEBERG TERMINAL", "Financial statement export"),
        ("Symbol", symbol),
        ("Exported UTC", datetime.now(timezone.utc).isoformat()),
        ("Units", "Currency values in millions; share counts in millions; per-share values unscaled"),
    ]
    for row in metadata_rows:
        metadata.append(row)
    _style_metadata(metadata)

    sheet_names: list[str] = []
    skipped: list[str] = []
    labels = {
        StatementType.INCOME: "Income",
        StatementType.BALANCE_SHEET: "Balance",
        StatementType.CASH_FLOW: "Cash Flow",
    }
    for statement in statements:
        for frequency, periods in (("Annual", statement.annual), ("Quarterly", statement.quarterly)):
            name = f"{labels[statement.statement_type]} {frequency}"
            if not periods:
                skipped.append(name)
                continue
            worksheet = workbook.create_sheet(name)
            _write_statement_sheet(worksheet, statement, periods, frequency)
            sheet_names.append(name)

    workbook.save(path)
    return FinancialExportResult(symbol, path, tuple(sheet_names), tuple(skipped))


def _write_statement_sheet(
    worksheet,
    statement: FinancialStatements,
    periods: list[FinancialPeriod],
    frequency: str,
) -> None:
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "C11"
    worksheet.append([f"{statement.name} ({statement.symbol})"])
    worksheet.append([f"{_statement_title(statement.statement_type)} / {frequency.upper()}"])
    worksheet.append(
        [f"Currency: {statement.currency or '--'} | Units: millions except per-share values"]
    )
    worksheet.append([f"Provider: {statement.provider} | Quality: {statement.quality}"])
    worksheet.append([f"Source: {statement.source_url or statement.provider}"])
    headers = ["Metric", "Unit", *[_period_label(period, frequency) for period in periods]]
    worksheet.append(headers)
    worksheet.append(
        [
            "Period End",
            "",
            *[period.end_date if period.end_date else None for period in periods],
        ]
    )
    worksheet.append(
        [
            "Source Form",
            "",
            *[
                f"{period.source_form}{' / DERIVED' if period.derived else ''}" or "--"
                for period in periods
            ],
        ]
    )
    worksheet.append(
        [
            "Filed Date",
            "",
            *[period.filed_date if period.filed_date else None for period in periods],
        ]
    )
    worksheet.append(
        [
            "SEC Accession",
            "",
            *[period.accession_number or "--" for period in periods],
        ]
    )

    metrics = ordered_statement_metrics(
        periods,
        statement.statement_type,
        statement.metric_order,
    )
    for metric in metrics:
        unit, scale, number_format = _metric_display(metric, statement.currency)
        is_expense = is_expense_or_outflow_metric(metric, statement.statement_type)
        worksheet.append(
            [
                statement.metric_labels.get(metric) or _metric_label(metric),
                unit,
                *[_scaled_value(period.values.get(metric), scale) for period in periods],
            ]
        )
        current_row = worksheet[worksheet.max_row]
        if is_expense:
            current_row[0].font = Font(color="C00000")
        for cell in current_row[2:]:
            cell.number_format = number_format
            cell.alignment = Alignment(horizontal="right")
            if is_expense:
                cell.font = Font(color="C00000")

    accent = PatternFill("solid", fgColor="F2A900")
    dark = PatternFill("solid", fgColor="171717")
    for cell in worksheet[1]:
        cell.fill = dark
        cell.font = Font(color="FFFFFF", bold=True, size=13)
    for cell in worksheet[6]:
        cell.fill = accent
        cell.font = Font(color="000000", bold=True)
        cell.alignment = Alignment(horizontal="right" if cell.column > 2 else "left")
    for row_number in (7, 8, 9, 10):
        for cell in worksheet[row_number][2:]:
            if row_number in (7, 9):
                cell.number_format = "dd-mmm-yyyy"
            cell.font = Font(color="666666", italic=True)
            cell.alignment = Alignment(horizontal="right")
    worksheet.column_dimensions["A"].width = 38
    worksheet.column_dimensions["B"].width = 14
    for index in range(3, len(headers) + 1):
        worksheet.column_dimensions[get_column_letter(index)].width = 15
    worksheet.auto_filter.ref = f"A6:{get_column_letter(len(headers))}{worksheet.max_row}"


def _style_metadata(worksheet) -> None:
    worksheet.sheet_view.showGridLines = False
    worksheet.column_dimensions["A"].width = 24
    worksheet.column_dimensions["B"].width = 62
    for cell in worksheet[1]:
        cell.fill = PatternFill("solid", fgColor="171717")
        cell.font = Font(color="F2A900", bold=True)
    for row in worksheet.iter_rows(min_row=2, max_col=1):
        row[0].font = Font(bold=True)


def _period_label(period: FinancialPeriod, frequency: str) -> str:
    label = period.period.strip().upper()
    year = period.end_date.year if period.end_date else None
    if frequency.upper() == "ANNUAL":
        fiscal_year = re.search(r"(?:FY\s*)?(\d{4})", label)
        return f"FY {fiscal_year.group(1) if fiscal_year else year or label}"
    quarter = re.search(r"Q([1-4])(?:\s+(\d{4}))?", label)
    if quarter:
        return f"Q{quarter.group(1)} {quarter.group(2) or year or ''}".strip()
    return label or (str(year) if year else "--")


def _statement_title(statement_type: StatementType) -> str:
    return {
        StatementType.INCOME: "INCOME STATEMENT",
        StatementType.BALANCE_SHEET: "BALANCE SHEET",
        StatementType.CASH_FLOW: "CASH FLOW STATEMENT",
    }[statement_type]


def _metric_display(metric: str, currency: str) -> tuple[str, float, str]:
    normalized = re.sub(r"[^a-z0-9]", "", metric.lower())
    money = currency or "CCY"
    if "eps" in normalized or "pershare" in normalized:
        return f"{money} / share", 1.0, '#,##0.00;[Red](#,##0.00);--'
    if "shares" in normalized or normalized.endswith("sharesnumber"):
        return "M shares", 1_000_000.0, '#,##0.0;[Red](#,##0.0);--'
    if any(token in normalized for token in ("margin", "rate", "ratio")):
        return "%", 1.0, '0.0%;[Red](0.0%);--'
    return f"{money} M", 1_000_000.0, '#,##0.0;[Red](#,##0.0);--'


def _scaled_value(value: float | None, scale: float) -> float | None:
    return value / scale if value is not None else None


_METRIC_LABELS = {
    "totalrevenue": "Revenue",
    "costofrevenue": "Cost Of Revenue",
    "grossprofit": "Gross Profit",
    "operatingexpense": "Operating Expenses",
    "operatingincome": "Operating Income",
    "normalizedebitda": "Normalized EBITDA",
    "pretaxincome": "Pre-Tax Income",
    "taxprovision": "Income Tax Expense",
    "netincome": "Net Income",
    "netincomecommonstockholders": "Net Income To Common",
    "basiceps": "Basic EPS",
    "dilutedeps": "Diluted EPS",
    "basicaverageshares": "Basic Average Shares",
    "dilutedaverageshares": "Diluted Average Shares",
    "interestexpense": "Interest Expense",
    "researchanddevelopment": "Research & Development",
    "sellinggeneralandadministration": "Selling, General & Administrative",
    "totalassets": "Total Assets",
    "currentassets": "Current Assets",
    "cashcashequivalentsandshortterminvestments": "Cash & Short-Term Investments",
    "cashandcashequivalents": "Cash & Cash Equivalents",
    "othershortterminvestments": "Other Short-Term Investments",
    "totalnoncurrentassets": "Total Non-Current Assets",
    "netppe": "Net Property, Plant & Equipment",
    "goodwillandotherintangibleassets": "Goodwill & Intangible Assets",
    "totalliabilitiesnetminorityinterest": "Total Liabilities",
    "currentliabilities": "Current Liabilities",
    "totalnoncurrentliabilitiesnetminorityinterest": "Non-Current Liabilities",
    "totaldebt": "Total Debt",
    "currentdebt": "Current Debt",
    "longtermdebt": "Long-Term Debt",
    "stockholdersequity": "Stockholders' Equity",
    "commonstockequity": "Common Equity",
    "retainedearnings": "Retained Earnings",
    "workingcapital": "Working Capital",
    "netdebt": "Net Debt",
    "ordinarysharesnumber": "Ordinary Shares Outstanding",
    "operatingcashflow": "Operating Cash Flow",
    "capitalexpenditure": "Capital Expenditure",
    "freecashflow": "Free Cash Flow",
    "investingcashflow": "Investing Cash Flow",
    "financingcashflow": "Financing Cash Flow",
    "endcashposition": "Ending Cash Position",
    "changesincash": "Change In Cash",
    "repurchaseofcapitalstock": "Share Repurchases",
    "cashdividendspaid": "Cash Dividends Paid",
    "issuanceofdebt": "Debt Issuance",
    "repaymentofdebt": "Debt Repayment",
    "stockbasedcompensation": "Stock-Based Compensation",
    "depreciationandamortization": "Depreciation & Amortization",
    "freecashflowcalculated": "Free Cash Flow (Calculated)",
}


def _metric_label(metric: str) -> str:
    normalized = re.sub(r"[^a-z0-9]", "", metric.lower())
    if normalized in _METRIC_LABELS:
        return _METRIC_LABELS[normalized]
    spaced = re.sub(r"(?<=[a-z0-9])(?=[A-Z])", " ", metric.replace("_", " "))
    return " ".join(spaced.split()).title()
